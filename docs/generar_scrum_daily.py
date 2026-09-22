# -*- coding: utf-8 -*-
"""
Genera la tabla del Scrum Daily del Sprint 1 (SDGE-26).

Lo que se llena aquí es solo lo que se puede comprobar:
  - los días hábiles del sprint y quién estaba en qué historia, según las
    fechas de inicio y vencimiento de Jira;
  - el burndown, con la fecha real en que cada historia pasó a Finalizada
    en Jira (resolutiondate, consultada el 22/09/2026);
  - los impedimentos que dejaron rastro en Git, en Jira o en el README.

Lo que cada integrante hizo y dijo en la reunión (qué hice, qué haré,
impedimento del día) lo escribe el equipo: no se inventa.
"""

import os
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, Reference
from openpyxl.worksheet.datavalidation import DataValidation

INICIO, FIN, HOY = date(2026, 9, 9), date(2026, 9, 28), date(2026, 9, 22)
DIAS = [INICIO + timedelta(d) for d in range((FIN - INICIO).days + 1)
        if (INICIO + timedelta(d)).weekday() < 5]

# clave, titulo, responsable, SP, inicio, vencimiento, finalizada (Jira)
HISTORIAS = [
    ("SDGE-1",  "Interfaz de inicio de sesión",                 "Eduardo", 3, "09-09", "09-10", "09-12"),
    ("SDGE-2",  "Maquetación de la página de inicio",           "Eder",    5, "09-09", "09-16", "09-18"),
    ("SDGE-3",  "Encabezado con navegación y botones de sesión","Eder",    3, "09-09", "09-10", "09-12"),
    ("SDGE-4",  "Sección hero con buscador de eventos",         "Eder",    5, "09-11", "09-21", "09-22"),
    ("SDGE-6",  "Interfaz de registro de administrador",        "Samuel",  3, "09-09", "09-10", "09-12"),
    ("SDGE-7",  "Interfaz de registro de cliente",              "Eduardo", 5, "09-11", "09-16", "09-18"),
    ("SDGE-8",  "Validación de formularios en el navegador",    "Samuel",  5, "09-11", "09-21", "09-22"),
    ("SDGE-9",  "Interfaz de registro de agente",               "Eduardo", 5, "09-17", "09-21", "09-22"),
    ("SDGE-10", "Ajuste responsive de las interfaces de registro","Jorge", 5, "09-17", "09-25", None),
    ("SDGE-11", "Sección de eventos destacados y categorías",   "Samuel",  5, "09-11", "09-16", "09-18"),
    ("SDGE-12", "Hoja de estilos del sistema de diseño",        "Jorge",   5, "09-09", "09-11", "09-12"),
    ("SDGE-14", "Maquetación Mobile-First de la página de inicio","Jorge", 3, "09-22", "09-22", "09-22"),
    ("SDGE-21", "Consolidar la estructura del repositorio en GitHub","Eder",1, "09-22", "09-22", None),
    ("SDGE-26", "Documentar el Scrum Daily del Sprint 1",       "Eduardo", 3, "09-09", "09-28", None),
]
# Las tres últimas entraron al sprint el 22/09 (cambio de alcance).
AGREGADAS_22 = {"SDGE-14", "SDGE-21", "SDGE-26"}

INTEGRANTES = [("Eder", "Eder Fabián Rodríguez Murillo"),
               ("Eduardo", "Eduardo José Benítez Guevara"),
               ("Jorge", "Jorge Andrés Marín Díaz"),
               ("Samuel", "Samuel Uribe Naranjo")]

IMPEDIMENTOS = [
    ("Sprint 1", "Samuel",
     "Samuel no aparecía en el proyecto de Jira, así que sus historias no tenían responsable.",
     "Se le invitó al proyecto y se repartió el sprint entre los cuatro integrantes.",
     "Resuelto", "Jira: SDGE-6, SDGE-8 y SDGE-11 asignadas a Samuel"),
    ("Sprint 1", "Samuel",
     "La pantalla de registro de administrador no existía en el diseño de Figma.",
     "Se diseñó sobre el mismo sistema de diseño, con código de invitación y correo institucional.",
     "Resuelto", "Jira SDGE-6 · README, nota sobre el registro de administrador"),
    ("12/09/2026", "Eder",
     "GitHub rechazó el primer push por falta de permisos en la cuenta.",
     "Se corrigió el acceso de la cuenta al repositorio y se subió el Sprint 1.",
     "Resuelto", "Git: commit c2332cc"),
    ("12/09/2026", "Todos",
     "Retroalimentación del profesor: los colores se confundían y sobraba espacio en blanco.",
     "Se unificó el color de las categorías y se redujo la separación entre secciones de 80 a 56 px.",
     "Resuelto", "Git: commits 65a306a y cbb30c9"),
    ("22/09/2026", "Jorge",
     "La hoja de estilos estaba escrita para escritorio (max-width) y no cumplía con Mobile-First.",
     "Se creó SDGE-14 y se reescribió con estilos base de celular y min-width.",
     "Resuelto", "Jira SDGE-14 · Git: commit 111cd85"),
    ("22/09/2026", "Jorge",
     "En el menú de celular, los enlaces y los botones se abrían uno encima del otro.",
     "El menú ahora se despliega en el flujo, debajo del logo.",
     "Resuelto", "Git: commit 111cd85"),
    ("22/09/2026", "Samuel",
     "SDGE-11 tenía la fecha de inicio (17/09) después de la de vencimiento (16/09).",
     "Se corrigió a 11/09 – 16/09, como estaba en la planeación.",
     "Resuelto", "Jira SDGE-11"),
    ("22/09/2026", "Eder",
     "Los compañeros todavía no son colaboradores del repositorio en GitHub.",
     "Invitarlos desde Settings → Collaborators.",
     "Pendiente", "Jira SDGE-24"),
]

# --- estilo -----------------------------------------------------------------
AZUL, TINTA, GRIS = "493EE5", "191C1E", "5A5868"
fino = Side(style="thin", color="D9D8E0")
BORDE = Border(left=fino, right=fino, top=fino, bottom=fino)
CAB_F = PatternFill("solid", fgColor="EDECFC")
ZEBRA = PatternFill("solid", fgColor="F7F7FA")
LLENAR = PatternFill("solid", fgColor="FFF8E6")   # celdas que llena el equipo
AJ = Alignment(wrap_text=True, vertical="top")
CEN = Alignment(horizontal="center", vertical="top")


def d(mmdd):
    return date(2026, int(mmdd[:2]), int(mmdd[3:]))


def cabecera(ws, fila, textos, anchos):
    for i, (t, w) in enumerate(zip(textos, anchos), start=1):
        c = ws.cell(row=fila, column=i, value=t)
        c.font = Font(bold=True, color=AZUL, size=10)
        c.fill = CAB_F
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDE
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[fila].height = 32


def titulo(ws, texto, sub):
    ws["A1"] = texto
    ws["A1"].font = Font(bold=True, size=15, color=TINTA)
    ws["A2"] = sub
    ws["A2"].font = Font(italic=True, size=10, color=GRIS)


wb = Workbook()

# ===========================================================================
# Hoja 2 primero (la 1 la referencia): Burndown
# ===========================================================================
bd = wb.active
bd.title = "Burndown"
titulo(bd, "Indicador de avance · Burndown del Sprint 1",
       "Puntos pendientes al cierre de cada día hábil. Los puntos terminados salen de la fecha "
       "en que cada historia pasó a Finalizada en Jira. El 22/09 el alcance subió de 49 a 56 SP "
       "(SDGE-14, SDGE-21 y SDGE-26).")
cabecera(bd, 4, ["Día", "Fecha", "Alcance (SP)", "Terminados ese día (SP)",
                 "Terminados acumulados (SP)", "Pendientes reales (SP)",
                 "Pendientes ideales (SP)", "Avance del sprint"],
         [7, 12, 12, 14, 14, 14, 14, 12])

n = len(DIAS)
for k, dia in enumerate(DIAS):
    f = 5 + k
    alcance = sum(h[3] for h in HISTORIAS if h[0] not in AGREGADAS_22 or dia >= date(2026, 9, 22))
    # lo terminado en fin de semana se cuenta el siguiente día hábil
    previo = DIAS[k - 1] if k else INICIO - timedelta(1)
    hechos = sum(h[3] for h in HISTORIAS if h[6] and previo < d(h[6]) <= dia)
    bd.cell(row=f, column=1, value=k + 1).alignment = CEN
    c = bd.cell(row=f, column=2, value=dia); c.number_format = "DD/MM/YYYY"; c.alignment = CEN
    bd.cell(row=f, column=3, value=alcance).alignment = CEN
    if dia <= HOY:
        bd.cell(row=f, column=4, value=hechos).alignment = CEN
        bd.cell(row=f, column=5, value="=SUM($D$5:D%d)" % f).alignment = CEN
        bd.cell(row=f, column=6, value="=C%d-E%d" % (f, f)).alignment = CEN
        c = bd.cell(row=f, column=8, value="=E%d/C%d" % (f, f)); c.number_format = "0%"; c.alignment = CEN
    else:
        for col in (4, 5, 6, 8):
            bd.cell(row=f, column=col).fill = LLENAR
    # línea ideal: del alcance inicial a 0 en el último día
    bd.cell(row=f, column=7, value="=ROUND($C$5*(1-(A%d-1)/%d),1)" % (f, n - 1)).alignment = CEN
    for col in range(1, 9):
        bd.cell(row=f, column=col).border = BORDE
        if k % 2:
            if bd.cell(row=f, column=col).fill != LLENAR:
                bd.cell(row=f, column=col).fill = ZEBRA

ultima = 4 + n
bd.cell(row=ultima + 2, column=1,
        value="Celdas amarillas: días que aún no han pasado. Se llenan con los puntos que se "
              "terminen ese día.").font = Font(italic=True, size=9, color=GRIS)

graf = LineChart()
graf.title = "Burndown · Sprint 1"
graf.y_axis.title = "Puntos pendientes"
graf.x_axis.title = "Día hábil"
graf.height, graf.width = 9, 18
datos = Reference(bd, min_col=6, max_col=7, min_row=4, max_row=ultima)
graf.add_data(datos, titles_from_data=True)
graf.set_categories(Reference(bd, min_col=1, min_row=5, max_row=ultima))
graf.series[0].graphicalProperties.line.solidFill = AZUL
graf.series[0].graphicalProperties.line.width = 28000
graf.series[1].graphicalProperties.line.solidFill = "A8A5B5"
graf.series[1].graphicalProperties.line.dashStyle = "dash"
bd.add_chart(graf, "J4")

# ===========================================================================
# Hoja 1: Daily
# ===========================================================================
dl = wb.create_sheet("Scrum Daily", 0)
titulo(dl, "Scrum Daily · Sprint 1 · TicketFlow (SDGE)",
       "9 – 28 de septiembre de 2026 · Reunión diaria de 15 minutos. Columnas blancas: salen de Jira. "
       "Columnas amarillas: las llena cada integrante con lo que dijo en la reunión.")
cabecera(dl, 4, ["Fecha", "Integrante", "Historia en curso (Jira)", "¿Qué hice ayer?",
                 "¿Qué haré hoy?", "Impedimento", "Solución", "Avance del sprint"],
         [12, 12, 38, 34, 34, 30, 30, 11])

fila = 5
for k, dia in enumerate(DIAS):
    fb = 5 + k   # fila del mismo día en Burndown
    for corto, _ in INTEGRANTES:
        activas = [h for h in HISTORIAS
                   if h[2] == corto and d(h[4]) <= dia <= d(h[5])
                   and (h[0] not in AGREGADAS_22 or dia >= date(2026, 9, 22))]
        texto = "\n".join("%s · %s" % (h[0], h[1]) for h in activas) or "Sin historia asignada"
        c = dl.cell(row=fila, column=1, value=dia); c.number_format = "DD/MM/YYYY"
        dl.cell(row=fila, column=2, value=corto)
        dl.cell(row=fila, column=3, value=texto)
        for col in (4, 5, 6, 7):
            dl.cell(row=fila, column=col).fill = LLENAR
        if dia <= HOY:
            c = dl.cell(row=fila, column=8, value="=Burndown!H%d" % fb)
            c.number_format = "0%"
        for col in range(1, 9):
            cel = dl.cell(row=fila, column=col)
            cel.border = BORDE
            cel.alignment = CEN if col in (1, 2, 8) else AJ
            cel.font = Font(size=10, color=TINTA if col != 3 else GRIS)
        dl.row_dimensions[fila].height = max(30, 15 * texto.count("\n") + 30)
        fila += 1
    # separador visual entre días
    for col in range(1, 9):
        dl.cell(row=fila - 1, column=col).border = Border(left=fino, right=fino, top=fino,
                                                          bottom=Side(style="medium", color=AZUL))
dl.freeze_panes = "C5"
dl.auto_filter.ref = "A4:H%d" % (fila - 1)

# ===========================================================================
# Hoja 3: Impedimentos
# ===========================================================================
im = wb.create_sheet("Impedimentos")
titulo(im, "Impedimentos del Sprint 1 y su solución",
       "Solo los que dejaron rastro verificable en Git, Jira o el README. Agreguen debajo los que "
       "salieron en las reuniones.")
cabecera(im, 4, ["Fecha", "Integrante", "Impedimento", "Solución", "Estado", "Evidencia"],
         [12, 11, 46, 46, 11, 32])
dv = DataValidation(type="list", formula1='"Resuelto,En curso,Pendiente"', allow_blank=True)
im.add_data_validation(dv)
for k, row in enumerate(IMPEDIMENTOS):
    f = 5 + k
    for col, v in enumerate(row, start=1):
        c = im.cell(row=f, column=col, value=v)
        c.border = BORDE
        c.alignment = CEN if col in (1, 2, 5) else AJ
        c.font = Font(size=10, bold=(col == 5),
                      color=("16A34A" if v == "Resuelto" else "B45309") if col == 5 else TINTA)
    im.row_dimensions[f].height = 44
    dv.add(im.cell(row=f, column=5))
for f in range(5 + len(IMPEDIMENTOS), 5 + len(IMPEDIMENTOS) + 8):   # filas libres
    for col in range(1, 7):
        im.cell(row=f, column=col).border = BORDE
        im.cell(row=f, column=col).fill = LLENAR
    dv.add(im.cell(row=f, column=5))

# ===========================================================================
# Hoja 4: Resumen
# ===========================================================================
rs = wb.create_sheet("Resumen")
titulo(rs, "Resumen del Sprint 1", "Se actualiza solo a partir del Burndown.")
cab = ["Indicador", "Valor"]
cabecera(rs, 4, cab, [40, 16])
ind = [
    ("Puntos comprometidos (alcance final)", "=Burndown!C%d" % ultima),
    ("Puntos terminados a hoy (22/09)", "=Burndown!E%d" % (5 + DIAS.index(HOY))),
    ("Puntos pendientes a hoy", "=Burndown!F%d" % (5 + DIAS.index(HOY))),
    ("Avance del sprint a hoy", "=Burndown!H%d" % (5 + DIAS.index(HOY))),
    ("Historias y tareas en el sprint", len(HISTORIAS)),
    ("Historias y tareas terminadas", sum(1 for h in HISTORIAS if h[6])),
    ("Impedimentos resueltos", "=COUNTIF(Impedimentos!E:E,\"Resuelto\")"),
    ("Impedimentos pendientes", "=COUNTIF(Impedimentos!E:E,\"Pendiente\")+COUNTIF(Impedimentos!E:E,\"En curso\")"),
]
for k, (a, b) in enumerate(ind):
    f = 5 + k
    rs.cell(row=f, column=1, value=a).border = BORDE
    c = rs.cell(row=f, column=2, value=b); c.border = BORDE; c.alignment = CEN
    c.font = Font(bold=True, color=AZUL, size=11)
    if "Avance" in a:
        c.number_format = "0%"

rs.cell(row=15, column=1, value="Carga por integrante (SP)").font = Font(bold=True, color=TINTA)
cabecera(rs, 16, ["Integrante", "SP asignados"], [40, 16])
for k, (corto, largo) in enumerate(INTEGRANTES):
    f = 17 + k
    rs.cell(row=f, column=1, value=largo).border = BORDE
    c = rs.cell(row=f, column=2, value=sum(h[3] for h in HISTORIAS if h[2] == corto))
    c.border = BORDE; c.alignment = CEN

os.makedirs("entregables", exist_ok=True)
ruta = os.path.join("entregables", "Scrum_Daily_Sprint1_TicketFlow.xlsx")
wb.save(ruta)
print("Listo:", ruta, "| días hábiles:", len(DIAS), "| filas daily:", fila - 5)
