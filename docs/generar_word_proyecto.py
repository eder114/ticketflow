# -*- coding: utf-8 -*-
"""
Genera el documento de Word del Proyecto Integrador.

Contiene portada, tabla de contenido, introduccion, planteamiento, objetivos,
justificacion, marco teorico y legal, metodologia, resultados (PESTEL, Vester
y Canvas), conclusiones y bibliografia.

Las referencias legales de la bibliografia estan verificadas contra la fuente
oficial; la Resolucion 38063 de 2026 es un pliego de cargos, no una sancion
en firme, y asi se dice en el texto.
"""

import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from vester import PROBLEMAS, M, CLAVES, NOMBRE, calcular, clasificar

INF, DEP = calcular()
CLS, CX, CY = clasificar(INF, DEP)

AZUL  = RGBColor(0x49, 0x3E, 0xE5)
TINTA = RGBColor(0x19, 0x1C, 0x1E)
GRIS  = RGBColor(0x5A, 0x58, 0x68)

doc = Document()

# ---------------------------------------------------------------- estilos --
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.font.color.rgb = TINTA
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
pf = normal.paragraph_format
pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
pf.space_after = Pt(8)
pf.line_spacing = 1.25

for nivel, tam in ((1, 16), (2, 13), (3, 11.5)):
    st = doc.styles["Heading %d" % nivel]
    st.font.name = "Cambria"
    st.font.size = Pt(tam)
    st.font.bold = True
    st.font.color.rgb = TINTA if nivel == 1 else AZUL
    st.paragraph_format.space_before = Pt(18 if nivel == 1 else 12)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.keep_with_next = True

for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(2.5)
    s.left_margin = s.right_margin = Cm(2.8)


# --------------------------------------------------------------- ayudantes --
def p(texto="", estilo=None, align=None, tam=None, negrita=False,
      cursiva=False, color=None, antes=None, despues=None, sangria=None):
    par = doc.add_paragraph(style=estilo)
    if texto:
        r = par.add_run(texto)
        r.bold = negrita
        r.italic = cursiva
        if tam:
            r.font.size = Pt(tam)
        if color:
            r.font.color.rgb = color
    if align is not None:
        par.alignment = align
    if antes is not None:
        par.paragraph_format.space_before = Pt(antes)
    if despues is not None:
        par.paragraph_format.space_after = Pt(despues)
    if sangria is not None:
        par.paragraph_format.left_indent = Cm(sangria)
    return par


def mixto(partes, align=WD_ALIGN_PARAGRAPH.JUSTIFY, sangria=None):
    """Un parrafo con tramos en negrita: [(texto, negrita), ...]"""
    par = doc.add_paragraph()
    par.alignment = align
    if sangria is not None:
        par.paragraph_format.left_indent = Cm(sangria)
    for txt, neg in partes:
        r = par.add_run(txt)
        r.bold = neg
    return par


def vineta(texto, negrita_inicial=""):
    par = doc.add_paragraph(style="List Bullet")
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    par.paragraph_format.space_after = Pt(4)
    if negrita_inicial:
        par.add_run(negrita_inicial).bold = True
    par.add_run(texto)
    return par


def sombrear(celda, hexcol):
    tc = celda._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:fill"), hexcol)
    tc.append(sh)


def tabla(cabeceras, filas, anchos=None, tam=9.5, centrar_desde=None):
    t = doc.add_table(rows=1, cols=len(cabeceras))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(cabeceras):
        c = t.rows[0].cells[i]
        c.text = ""
        par = c.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = par.add_run(h)
        r.bold = True
        r.font.size = Pt(tam)
        r.font.color.rgb = GRIS
        sombrear(c, "F3F3F7")
    for fila in filas:
        celdas = t.add_row().cells
        for i, v in enumerate(fila):
            celdas[i].text = ""
            par = celdas[i].paragraphs[0]
            par.paragraph_format.space_after = Pt(2)
            if centrar_desde is not None and i >= centrar_desde:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = par.add_run(str(v))
            r.font.size = Pt(tam)
    if anchos:
        for fila in t.rows:
            for i, a in enumerate(anchos):
                fila.cells[i].width = Cm(a)
    return t


def pie_de_figura(texto):
    par = p(texto, align=WD_ALIGN_PARAGRAPH.CENTER, tam=9, cursiva=True,
            color=GRIS, antes=4, despues=14)
    return par


def salto_pagina():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# =========================================================================
# PORTADA
# =========================================================================
for _ in range(4):
    doc.add_paragraph()

p("TICKETFLOW", align=WD_ALIGN_PARAGRAPH.CENTER, tam=30, negrita=True, color=AZUL,
  despues=2)
p("Plataforma web para la gestión y reserva de boletería de eventos en Colombia",
  align=WD_ALIGN_PARAGRAPH.CENTER, tam=14, color=GRIS, despues=28)

p("Proyecto Integrador", align=WD_ALIGN_PARAGRAPH.CENTER, tam=13, negrita=True,
  despues=2)
p("Análisis del problema mediante Matriz de Vester y formulación del modelo de "
  "negocio con el lienzo Canvas", align=WD_ALIGN_PARAGRAPH.CENTER, tam=11,
  color=GRIS, despues=36)

p("Presentado por", align=WD_ALIGN_PARAGRAPH.CENTER, tam=10, color=GRIS, despues=6)
for nombre in ("Alex Andrés Cruz Rueda", "Eduardo José Benítez",
               "Jorge Andrés Marín", "Eder Fabián Rodríguez Murillo",
               "Samuel Uribe Naranjo"):
    p(nombre, align=WD_ALIGN_PARAGRAPH.CENTER, tam=11.5, despues=2)

doc.add_paragraph()
p("Unidad Central del Valle del Cauca — UCEVA", align=WD_ALIGN_PARAGRAPH.CENTER,
  tam=11.5, negrita=True, antes=24, despues=2)
p("Facultad de Ingeniería", align=WD_ALIGN_PARAGRAPH.CENTER, tam=11, color=GRIS,
  despues=2)
p("Tuluá, Valle del Cauca", align=WD_ALIGN_PARAGRAPH.CENTER, tam=11, color=GRIS,
  despues=2)
p("Septiembre de 2026", align=WD_ALIGN_PARAGRAPH.CENTER, tam=11, color=GRIS)

salto_pagina()

# =========================================================================
# TABLA DE CONTENIDO
# =========================================================================
p("Tabla de contenido", tam=16, negrita=True, align=WD_ALIGN_PARAGRAPH.LEFT,
  despues=12)
par = doc.add_paragraph()
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), r'TOC \o "1-2" \h \z \u')
sub = OxmlElement("w:r")
t_ = OxmlElement("w:t")
t_.text = "Haga clic aquí y pulse F9 para actualizar la tabla de contenido."
sub.append(t_)
fld.append(sub)
par._p.append(fld)

salto_pagina()

# =========================================================================
# 1. INTRODUCCION
# =========================================================================
doc.add_heading("1. Introducción", level=1)

p("La venta de boletería para eventos en Colombia mueve conciertos, festivales, "
  "partidos, funciones de teatro y salidas turísticas, pero lo hace sobre una "
  "infraestructura desigual. En un extremo están las plataformas grandes, que "
  "concentran los eventos masivos y cuyas prácticas comerciales ya han sido "
  "cuestionadas por la autoridad de protección al consumidor. En el otro extremo "
  "está una mayoría de organizadores pequeños —teatros municipales, salas "
  "independientes, operadores turísticos, promotores locales— que siguen llevando "
  "el aforo en una hoja de cálculo o en un cuaderno, confirmando reservas por "
  "WhatsApp y resolviendo las quejas por teléfono.")

p("Este documento presenta el análisis del problema que sustenta a TicketFlow, "
  "una plataforma web para la gestión y reserva de boletería. El análisis se hizo "
  "en dos capas complementarias. La primera es un análisis PESTEL, que mira hacia "
  "afuera y describe el entorno político, económico, social, tecnológico, ambiental "
  "y legal en el que la plataforma tendría que operar. La segunda es una matriz de "
  "Vester, que mira hacia adentro del problema y establece cuál de sus partes causa "
  "a cuál, para distinguir las raíces de las consecuencias.")

p("Esa distinción es la que le da sentido al documento. El PESTEL dice que el "
  "entorno es favorable, pero no dice por dónde empezar: enumera oportunidades y "
  "amenazas sin jerarquizarlas causalmente. La matriz de Vester sí ordena el "
  "problema, y al hacerlo obliga a revisar el modelo de negocio. Por eso el "
  "documento no termina en el diagnóstico: termina en un lienzo Canvas reformulado "
  "y en la comparación explícita con el Canvas que el equipo había construido antes, "
  "cuando solo contaba con el PESTEL.")

# =========================================================================
# 2. PLANTEAMIENTO
# =========================================================================
doc.add_heading("2. Planteamiento del problema", level=1)

doc.add_heading("2.1. Descripción", level=2)
p("El comprador de boletería en Colombia enfrenta tres fricciones reconocibles. "
  "La primera es de precio: el valor que aparece en la tarjeta del evento no es el "
  "que termina pagando, porque los cargos por servicio se revelan al final del "
  "proceso de compra. La segunda es de confianza: no tiene certeza de que su cupo "
  "quede efectivamente apartado, ni de que una cancelación se le devuelva en un "
  "plazo razonable. La tercera es de recurso: cuando algo sale mal, el canal de "
  "reclamación no deja rastro y no hay un radicado al cual referirse.")

p("Del lado del organizador el problema es distinto y menos visible. Un promotor "
  "pequeño no tiene acceso a una herramienta de gestión: administra el aforo a mano, "
  "no sabe en tiempo real cuántos cupos le quedan, y esa falta de control es lo que "
  "produce la sobreventa, las devoluciones que no alcanza a tramitar y la ausencia "
  "de un canal formal de quejas. Es decir, buena parte de lo que el comprador "
  "experimenta como mal servicio nace de que el organizador no tiene con qué "
  "prestarlo bien.")

p("Estas fricciones no son hipotéticas. En junio de 2026 la Superintendencia de "
  "Industria y Comercio formuló pliego de cargos contra la sociedad que opera "
  "Tuboleta.com por presuntas fallas en el deber de información, cláusulas abusivas "
  "y el incumplimiento de las reglas del comercio electrónico, incluidos el derecho "
  "de retracto y los mecanismos de PQRS. Se trata de una investigación en curso y no "
  "de una sanción en firme, pero muestra que los puntos de dolor identificados "
  "coinciden con los que la autoridad considera relevantes.")

doc.add_heading("2.2. Formulación", level=2)
p("¿Cuál de los problemas del mercado de boletería en Colombia es la causa raíz "
  "de los demás, y qué modelo de negocio se deriva de atacar esa causa en lugar de "
  "sus consecuencias?", cursiva=True)

# =========================================================================
# 3. OBJETIVOS
# =========================================================================
doc.add_heading("3. Objetivos", level=1)

doc.add_heading("3.1. Objetivo general", level=2)
p("Determinar, mediante una matriz de Vester, la estructura causal de los "
  "problemas del mercado de boletería en Colombia, y formular a partir de ella el "
  "modelo de negocio de TicketFlow en un lienzo Canvas.")

doc.add_heading("3.2. Objetivos específicos", level=2)
for o in (
    "Caracterizar el entorno del sector mediante un análisis PESTEL, identificando "
    "los factores de mayor incidencia sobre la operación de una plataforma de boletería.",
    "Identificar diez problemas del mercado a partir del enunciado del proyecto, "
    "del análisis PESTEL y del benchmark de las plataformas existentes.",
    "Construir la matriz de Vester cruzando esos diez problemas y clasificarlos en "
    "los cuadrantes activo, crítico, pasivo e indiferente.",
    "Formular el lienzo Canvas de TicketFlow tomando como insumo la clasificación "
    "obtenida, el análisis PESTEL y los requisitos del enunciado del proyecto.",
    "Contrastar el lienzo resultante con la versión previa del modelo de negocio y "
    "justificar cada cambio."
):
    vineta(o)

# =========================================================================
# 4. JUSTIFICACION
# =========================================================================
doc.add_heading("4. Justificación", level=1)

p("Un proyecto de software se justifica por el problema que resuelve, no por la "
  "tecnología que emplea. El riesgo habitual en un proyecto académico es construir "
  "una solución correcta para un problema secundario: atacar lo que se ve —la "
  "desconfianza del comprador, el abandono de la compra, la baja recompra— sin "
  "advertir que esas son consecuencias de algo que ocurre antes. La matriz de Vester "
  "existe precisamente para evitar ese error, porque obliga a declarar, casilla por "
  "casilla, qué causa qué.")

p("La justificación práctica es que el resultado de ese análisis es verificable y "
  "reorienta decisiones concretas: cambia cuál es el segmento principal de clientes, "
  "cuál es la fuente de ingresos que debe encabezar el modelo y qué funcionalidades "
  "son núcleo en lugar de accesorias. La justificación académica es que el ejercicio "
  "articula tres herramientas del seminario —PESTEL, Vester y Canvas— en una "
  "secuencia donde cada una alimenta a la siguiente, en lugar de presentarlas como "
  "entregables sueltos.")

# =========================================================================
# 5. MARCO REFERENCIAL
# =========================================================================
doc.add_heading("5. Marco referencial", level=1)

doc.add_heading("5.1. Marco teórico", level=2)

mixto([("Análisis PESTEL. ", True),
       ("Es una herramienta de exploración del entorno cuyo origen se atribuye a "
        "Francis J. Aguilar, quien en 1967 propuso examinar el ambiente de la empresa "
        "en categorías separadas. La formulación actual clasifica los factores externos "
        "en seis dimensiones: política, económica, social, tecnológica, ambiental y "
        "legal. Su función es descriptiva: dice qué hay alrededor del proyecto, pero no "
        "establece relaciones de causa y efecto entre esos factores.", False)])

mixto([("Matriz de Vester. ", True),
       ("Procede del pensamiento en red desarrollado por el bioquímico alemán Frederic "
        "Vester, miembro del Club de Roma, quien propuso analizar los sistemas complejos "
        "como una red de efectos interconectados en lugar de como una lista de elementos. "
        "En su aplicación habitual al diagnóstico de problemas, la matriz cruza cada "
        "problema contra todos los demás y valora, en una escala de 0 a 3, cuánto causa "
        "el problema de la fila al problema de la columna. La suma de la fila da la "
        "influencia y la suma de la columna, la dependencia. Al ubicar cada problema en "
        "un plano de influencia contra dependencia se obtienen cuatro cuadrantes:", False)])

for nom, desc in (
    ("Activos: ", "alta influencia y baja dependencia. Causan mucho y casi nada los "
     "causa; son las raíces del problema."),
    ("Críticos: ", "alta influencia y alta dependencia. Causan y son causados a la "
     "vez, por lo que el problema se realimenta a través de ellos."),
    ("Pasivos: ", "baja influencia y alta dependencia. Son consecuencias; no se "
     "atacan de frente."),
    ("Indiferentes: ", "baja influencia y baja dependencia. De prioridad menor, "
     "salvo que queden limítrofes con otro cuadrante."),
):
    vineta(desc, nom)

mixto([("Lienzo Canvas. ", True),
       ("Propuesto por Alexander Osterwalder e Yves Pigneur, describe un modelo de "
        "negocio en nueve bloques dispuestos sobre un lienzo: socios clave, actividades "
        "clave, recursos clave, propuesta de valor, relaciones con clientes, canales, "
        "segmentos de clientes, estructura de costos y fuentes de ingresos. Su valor en "
        "este trabajo es que, al ser una sola vista, obliga a que las decisiones sean "
        "coherentes entre sí: si cambia el segmento principal, tienen que cambiar con él "
        "la propuesta de valor, los canales y los ingresos.", False)])

doc.add_heading("5.2. Marco legal", level=2)

mixto([("Ley 1480 de 2011 (Estatuto del Consumidor). ", True),
       ("Establece el régimen general de protección al consumidor en Colombia, "
        "incluidos el deber de información, el derecho de retracto y la obligación de "
        "atender peticiones, quejas y reclamos.", False)])

mixto([("Ley 2439 de 2024. ", True),
       ("Modifica parcialmente el Estatuto del Consumidor en materia de comercio "
        "electrónico y entró en vigencia el 20 de diciembre de 2024. Su efecto más "
        "relevante para este proyecto es la reducción del plazo de devolución del dinero: "
        "cuando la operación se realiza por comercio electrónico, el reembolso debe "
        "completarse dentro de los quince días calendario siguientes al ejercicio del "
        "derecho de retracto, plazo que comprende la actuación de todos los "
        "intervinientes, incluidas las entidades bancarias. Además, el dinero debe "
        "devolverse por el mismo medio de pago que usó el consumidor. Antes de esta "
        "reforma el término era considerablemente más amplio.", False)])

mixto([("Ley 1581 de 2012. ", True),
       ("Régimen general de protección de datos personales, aplicable al tratamiento "
        "de los datos de clientes, agentes y administradores que la plataforma registra.",
        False)])

mixto([("Resolución 38063 de 2026 de la Superintendencia de Industria y Comercio. ", True),
       ("El 4 de junio de 2026 la SIC formuló pliego de cargos contra Ticket Fast "
        "S.A.S., operadora de Tuboleta.com, tras acumular nueve expedientes con quejas "
        "ciudadanas. Los cargos comprenden la omisión de información obligatoria sobre "
        "los eventos —municipio, promotor, horarios definitivos y canales de venta "
        "habilitados—, la no comunicación a la autoridad dentro de los tres días hábiles "
        "sobre el trámite de devoluciones en eventos cancelados o aplazados, la "
        "existencia de cláusulas abusivas, el incumplimiento de las reglas de comercio "
        "electrónico en cuanto al derecho de retracto y a los mecanismos de PQRS, y el "
        "desacato de órdenes previas de la autoridad. Debe subrayarse que se trata de un "
        "pliego de cargos: la investigación está en curso y la sociedad conserva su "
        "derecho de defensa. Se cita aquí no como declaración de responsabilidad, sino "
        "como evidencia de cuáles son los incumplimientos que la autoridad vigila en "
        "este mercado.", False)])

# =========================================================================
# 6. METODOLOGIA
# =========================================================================
doc.add_heading("6. Metodología", level=1)

p("El trabajo es de tipo descriptivo con enfoque mixto, y se desarrolló en cinco "
  "fases encadenadas, donde el producto de cada una es el insumo de la siguiente.")

for n, (tit, desc) in enumerate((
    ("Caracterización del entorno.",
     "Se aplicó el análisis PESTEL sobre el sector de boletería en Colombia. Cada "
     "factor se valoró con la fórmula Puntuación = Importancia × Intensidad × "
     "Tendencia, lo que permite ordenar las seis dimensiones por peso relativo en "
     "lugar de tratarlas como equivalentes."),
    ("Identificación de problemas.",
     "Se definieron diez problemas del mercado a partir de tres fuentes: el enunciado "
     "del proyecto integrador, los hallazgos del PESTEL y el benchmark de Tuboleta, "
     "Eventbrite, Ticketmaster y Taquilla Live. Se fijó el número en diez para que la "
     "matriz fuera manejable sin perder cobertura."),
    ("Construcción de la matriz de Vester.",
     "Se cruzó cada problema contra los nueve restantes con la escala 0 a 3. La "
     "diagonal principal se excluye, porque un problema no se causa a sí mismo. Como "
     "control de consistencia se verificó que la suma de todas las influencias iguale "
     "la suma de todas las dependencias, ya que ambas recorren las mismas casillas; el "
     "total obtenido fue 87 por ambos caminos."),
    ("Clasificación y lectura.",
     "Se calcularon la influencia y la dependencia de cada problema, se fijaron los "
     "cortes de los ejes en la mitad del rango de cada uno —8,5 en influencia y 10,0 en "
     "dependencia— y se ubicó cada problema en su cuadrante."),
    ("Formulación del modelo de negocio.",
     "Con la clasificación obtenida, el PESTEL y los requisitos del enunciado se "
     "diligenciaron los nueve bloques del lienzo Canvas, dejando trazabilidad de qué "
     "hallazgo respalda cada elemento, y se contrastó el resultado con la versión "
     "previa del modelo."),
), start=1):
    vineta(desc, "Fase %d — %s " % (n, tit))

p("El cálculo de la matriz se implementó en un script de Python que reproduce las "
  "sumas y la clasificación, y se replicó con fórmulas en el libro de Excel adjunto, "
  "de modo que ambos resultados puedan contrastarse de forma independiente.",
  antes=6)

# =========================================================================
# 7. RESULTADOS
# =========================================================================
doc.add_heading("7. Resultados", level=1)

doc.add_heading("7.1. Análisis PESTEL", level=2)
p("La valoración de los seis factores arrojó un puntaje total de 1.406 puntos, "
  "distribuidos así:")

tabla(["Factor", "Puntuación", "Participación"],
      [["Tecnológico", "305", "21,7 %"],
       ["Económico", "290", "20,6 %"],
       ["Social", "289", "20,6 %"],
       ["Político", "228", "16,2 %"],
       ["Legal", "210", "14,9 %"],
       ["Ambiental", "84", "6,0 %"],
       ["Total", "1.406", "100 %"]],
      anchos=[6.5, 4.0, 4.5], centrar_desde=1)
pie_de_figura("Tabla 1. Resultados del análisis PESTEL.")

p("Tres factores —tecnológico, económico y social— concentran cerca del 63 % del "
  "peso total, lo que indica que el entorno favorece la entrada de una plataforma "
  "digital: hay infraestructura disponible, capacidad de gasto en entretenimiento y "
  "un público habituado a comprar en línea. El factor legal, pese a ocupar el quinto "
  "lugar por puntaje, es el que concentra la amenaza más concreta, porque impone "
  "obligaciones exigibles con plazo: los quince días calendario para el reembolso y "
  "la trazabilidad de las PQRS no son buenas prácticas opcionales, sino condiciones "
  "de operación.")

doc.add_heading("7.2. Los diez problemas identificados", level=2)
tabla(["Código", "Problema", "Fuente"],
      [["P1", "Desconfianza del comprador en las plataformas de boletería",
        "Caso SIC; PESTEL legal"],
       ["P2", "Precio final oculto hasta el checkout", "Benchmark del sector"],
       ["P3", "Reventa y suplantación de boletas", "Benchmark del sector"],
       ["P4", "Sobreventa y descontrol del aforo por gestión manual",
        "Enunciado: Evento.capacidad"],
       ["P5", "Reembolsos lentos o incumplidos", "Ley 2439 de 2024; PESTEL legal"],
       ["P6", "Ausencia de un canal PQRS digital y trazable", "Caso SIC; PESTEL legal"],
       ["P7", "Organizadores pequeños sin herramientas digitales",
        "Enunciado: perfil Agente; PESTEL"],
       ["P8", "Caídas de la plataforma en picos de venta", "PESTEL tecnológico"],
       ["P9", "Abandono de la compra antes de confirmar", "Benchmark del sector"],
       ["P10", "Baja recompra y poca fidelización", "Enunciado: Cliente.puntos"]],
      anchos=[1.8, 8.4, 4.8], tam=9)
pie_de_figura("Tabla 2. Problemas que entran a la matriz.")

doc.add_heading("7.3. Matriz de Vester", level=2)
p("Cada casilla indica cuánto causa el problema de la fila al problema de la "
  "columna, en escala de 0 (no lo causa) a 3 (lo causa de forma directa y fuerte).")

filas_m = []
for i, cod in enumerate(CLAVES):
    fila = [cod] + [("—" if i == j else str(M[cod][j])) for j in range(10)] + [str(INF[cod])]
    filas_m.append(fila)
filas_m.append(["Dep."] + [str(DEP[c]) for c in CLAVES] + [str(sum(INF.values()))])

tabla(["", "P1", "P2", "P3", "P4", "P5", "P6", "P7", "P8", "P9", "P10", "Inf."],
      filas_m, anchos=[1.25] * 12, tam=8.5, centrar_desde=0)
pie_de_figura("Tabla 3. Matriz de Vester. La última columna es la influencia "
              "(suma de la fila) y la última fila, la dependencia (suma de la columna).")

p("La suma de las influencias y la suma de las dependencias coinciden en 87, lo que "
  "confirma que no hay errores de captura en el cruce.")

doc.add_heading("7.4. Clasificación por cuadrantes", level=2)
p("Con los cortes en 8,5 para la influencia y 10,0 para la dependencia, los diez "
  "problemas se distribuyen de la siguiente manera, ordenados de mayor a menor "
  "influencia:")

ETIQ = {"CRITICO": "Crítico", "ACTIVO": "Activo",
        "PASIVO": "Pasivo", "INDIFERENTE": "Indiferente"}
orden = sorted(CLAVES, key=lambda k: (-INF[k], -DEP[k]))
tabla(["Código", "Problema", "Influencia", "Dependencia", "Cuadrante"],
      [[c, NOMBRE[c].replace("(Ley 2439 de 2024)", "").replace("(cargos sorpresa)", "").strip(),
        INF[c], DEP[c], ETIQ[CLS[c]]] for c in orden],
      anchos=[1.7, 6.7, 2.2, 2.4, 2.4], tam=9, centrar_desde=2)
pie_de_figura("Tabla 4. Clasificación de los problemas.")

p("La lectura del resultado es la siguiente.")

mixto([("La causa raíz es P7, la ausencia de herramientas digitales en los "
        "organizadores pequeños.", True),
       (" Con 17 puntos de influencia y apenas 2 de dependencia, es el valor más alto "
        "de toda la matriz y el que menos depende de los demás. Un organizador que "
        "gestiona su aforo a mano termina produciendo sobreventa, no controla la "
        "reventa, no alcanza a tramitar los reembolsos y no tiene cómo operar un canal "
        "de PQRS. Cuatro problemas distintos con un mismo origen.", False)])

mixto([("El único problema crítico es P5, los reembolsos lentos o incumplidos.", True),
       (" Con 9 de influencia y 11 de dependencia queda en el cuadrante donde el "
        "problema se realimenta: es causado por la sobreventa y por la falta de un canal "
        "de quejas, y a su vez alimenta la desconfianza y la baja recompra. Es el punto "
        "donde cualquier mejora aguas arriba se pierde si no se resuelve, y coincide "
        "exactamente con la obligación legal de los quince días.", False)])

mixto([("Los problemas más visibles resultaron ser los menos importantes de atacar "
        "directamente.", True),
       (" La baja recompra (P10) tiene 20 de dependencia y 1 de influencia; la "
        "desconfianza del comprador (P1), 19 y 6; el abandono de la compra (P9), 14 y 3. "
        "Son consecuencias casi puras. Una estrategia dirigida a ellos —más publicidad, "
        "más descuentos, más programas de puntos— trataría el síntoma sin tocar la "
        "causa.", False)])

mixto([("P6 merece una observación.", True),
       (" Con 8 de influencia y 9 de dependencia queda formalmente como indiferente, "
        "pero los dos valores están apenas por debajo de sus respectivos cortes. Es un "
        "caso limítrofe con el cuadrante crítico y, dado que además corresponde a una "
        "obligación legal exigible, en la práctica conviene tratarlo como si lo fuera.",
        False)])

doc.add_heading("7.5. Modelo Canvas", level=2)
p("El lienzo resultante se presenta en la Figura 1. Cada elemento lleva la marca de "
  "su origen: Pn remite al problema correspondiente de la matriz de Vester, PESTEL a "
  "un factor del análisis del entorno, ENUNCIADO a un requisito explícito del "
  "enunciado del proyecto y EQUIPO a una decisión propia que no se desprende de "
  "ninguna de las tres fuentes anteriores. Esa trazabilidad es intencional: permite "
  "verificar qué está respaldado por una fuente y qué es criterio del equipo, en "
  "lugar de presentar todo el lienzo con el mismo grado de sustento.")

ruta_img = os.path.join("entregables", "Canvas_TicketFlow_v2.png")
if os.path.exists(ruta_img):
    # El lienzo es muy ancho: en una pagina vertical la letra quedaria de unos
    # 3 puntos. Va en una seccion horizontal propia, con margenes estrechos.
    apaisada = doc.add_section(WD_SECTION.NEW_PAGE)
    apaisada.orientation = WD_ORIENT.LANDSCAPE
    apaisada.page_width, apaisada.page_height = apaisada.page_height, apaisada.page_width
    apaisada.left_margin = apaisada.right_margin = Cm(1.5)
    apaisada.top_margin = apaisada.bottom_margin = Cm(1.8)

    ancho_util = apaisada.page_width - apaisada.left_margin - apaisada.right_margin
    doc.add_picture(ruta_img, width=ancho_util)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    pie_de_figura("Figura 1. Modelo Canvas de TicketFlow, versión 2. "
                  "El archivo Canvas_TicketFlow_v2.png acompaña este documento "
                  "en resolución completa.")

    # Se vuelve a vertical para el resto del documento.
    vertical = doc.add_section(WD_SECTION.NEW_PAGE)
    vertical.orientation = WD_ORIENT.PORTRAIT
    vertical.page_width, vertical.page_height = vertical.page_height, vertical.page_width
    vertical.top_margin = vertical.bottom_margin = Cm(2.5)
    vertical.left_margin = vertical.right_margin = Cm(2.8)

doc.add_heading("7.6. Contraste con el modelo anterior", level=2)
p("El equipo había construido una primera versión del Canvas cuando solo contaba con "
  "el análisis PESTEL. Al incorporar la matriz de Vester, cuatro bloques cambian de "
  "forma sustancial y dos de forma menor. Un lienzo que no hubiera cambiado sería "
  "señal de que la matriz no se usó.")

tabla(["Bloque", "Versión 1 (solo PESTEL)", "Versión 2 (con Vester)", "Razón del cambio"],
      [["Segmentos de clientes",
        "El cliente final encabezaba la lista; el organizador aparecía en segundo lugar.",
        "El organizador pasa a ser el segmento primario.",
        "P7 es el problema más activo de la matriz (17 de influencia). El cliente final "
        "es, en términos causales, casi pura consecuencia."],
       ["Propuesta de valor",
        "Una sola lista que mezclaba beneficios para ambos públicos.",
        "Dos promesas separadas: una para el organizador y otra para el cliente.",
        "Si son dos segmentos con problemas de naturaleza distinta, no pueden compartir "
        "una sola promesa."],
       ["Actividades clave",
        "Actividades genéricas de plataforma: desarrollo, soporte y mercadeo.",
        "Aparecen el control de aforo en tiempo real y el motor de reembolsos con plazo "
        "de quince días.",
        "Atacan P4 (activo) y P5 (el único crítico). En la versión anterior ninguna "
        "actividad los nombraba."],
       ["Fuentes de ingresos",
        "La comisión por transacción era el modelo principal.",
        "La suscripción del organizador pasa al primer lugar.",
        "Si el organizador es el segmento primario, el ingreso principal debe provenir "
        "de resolverle su problema."],
       ["Estructura de costos",
        "Costos habituales de operación y cumplimiento.",
        "Se incorpora la reserva operativa destinada a los reembolsos.",
        "Cumplir el plazo legal de quince días exige disponibilidad de caja, no "
        "solamente voluntad."],
       ["Recursos clave",
        "Plataforma, equipo, marca, convenios y base de datos.",
        "Se nombra explícitamente la máquina de estados de la reserva.",
        "Es el mecanismo concreto que impide la sobreventa; antes quedaba implícito "
        "dentro de «la plataforma»."]],
      anchos=[2.8, 4.0, 4.0, 4.5], tam=8.5)
pie_de_figura("Tabla 5. Cambios entre la versión 1 y la versión 2 del Canvas.")

# =========================================================================
# 8. CONCLUSIONES
# =========================================================================
doc.add_heading("8. Conclusiones", level=1)

for c in (
    "La matriz de Vester identificó como causa raíz del problema la ausencia de "
    "herramientas digitales en los organizadores pequeños (P7), con 17 puntos de "
    "influencia frente a 2 de dependencia. Este hallazgo valida el proyecto: el perfil "
    "de Agente previsto en el enunciado —que registra eventos con su capacidad total y "
    "administra las reservas actualizando su estado— es precisamente la herramienta "
    "cuya ausencia constituye ese problema. TicketFlow ataca la raíz y no una "
    "consecuencia.",

    "El único problema crítico es el incumplimiento en los reembolsos (P5), que causa y "
    "es causado a la vez. Su tratamiento no admite postergación, porque mientras persista "
    "neutraliza las mejoras que se logren aguas arriba. Su coincidencia con una obligación "
    "legal de plazo cierto lo convierte además en un requisito de operación y no en una "
    "mejora deseable.",

    "Los problemas de mayor visibilidad comercial —la desconfianza del comprador, el "
    "abandono de la compra y la baja recompra— resultaron ser consecuencias con muy poca "
    "capacidad de causar. Invertir en ellos directamente sería tratar el síntoma. Se "
    "resuelven por efecto de haber resuelto los activos y el crítico.",

    "El análisis PESTEL y la matriz de Vester responden preguntas distintas y ninguna "
    "sustituye a la otra. El PESTEL estableció que el entorno es favorable, con el 63 % "
    "del peso concentrado en los factores tecnológico, económico y social; la matriz "
    "estableció por dónde empezar. Es la combinación de ambas la que produce una decisión.",

    "La incorporación de la matriz obligó a modificar cuatro bloques del modelo de "
    "negocio, incluidos el segmento principal de clientes y la fuente de ingresos que "
    "encabeza el lienzo. Esto confirma que el análisis causal no fue un ejercicio "
    "documental, sino que cambió decisiones concretas del proyecto.",
):
    vineta(c)

# =========================================================================
# 9. BIBLIOGRAFIA
# =========================================================================
doc.add_heading("9. Bibliografía", level=1)

REFS = [
    "Aguilar, F. J. (1967). Scanning the Business Environment. Macmillan.",

    "Congreso de la República de Colombia. (2011, 12 de octubre). Ley 1480 de 2011, "
    "por medio de la cual se expide el Estatuto del Consumidor y se dictan otras "
    "disposiciones. Diario Oficial No. 48.220.",

    "Congreso de la República de Colombia. (2012, 17 de octubre). Ley 1581 de 2012, "
    "por la cual se dictan disposiciones generales para la protección de datos "
    "personales. Diario Oficial No. 48.587.",

    "Congreso de la República de Colombia. (2024). Ley 2439 de 2024, por medio de la "
    "cual se modifica parcialmente el Estatuto del Consumidor en materia de comercio "
    "electrónico. Vigente desde el 20 de diciembre de 2024. "
    "https://www.funcionpublica.gov.co/eva/gestornormativo/norma.php?i=257116",

    "Osterwalder, A. y Pigneur, Y. (2010). Business Model Generation: A Handbook for "
    "Visionaries, Game Changers, and Challengers. John Wiley & Sons.",

    "Superintendencia de Industria y Comercio. (2026, 4 de junio). La SIC del Cambio "
    "formula pliego de cargos a TuBoleta por presuntas fallas en el deber de "
    "información, las cláusulas abusivas y las reglas del comercio electrónico "
    "[Comunicado de prensa]. Resolución 38063 de 2026. "
    "https://sedeelectronica.sic.gov.co/comunicado/la-sic-del-cambio-formula-pliego-de-"
    "cargos-tuboleta-por-presuntas-fallas-en-el-deber-de-informacion-las-clausulas-"
    "abusivas-y-las-reglas",

    "Vester, F. (2002). Die Kunst, vernetzt zu denken: Ideen und Werkzeuge für einen "
    "neuen Umgang mit Komplexität [El arte de pensar en red]. Deutsche Verlags-Anstalt. "
    "Informe al Club de Roma.",
]
for r in REFS:
    par = p(r, despues=8)
    par.paragraph_format.left_indent = Cm(1.0)
    par.paragraph_format.first_line_indent = Cm(-1.0)

p("Los cálculos de la matriz se encuentran en el archivo "
  "Matriz_Vester_TicketFlow.xlsx, que acompaña a este documento, y el lienzo en "
  "Canvas_TicketFlow_v2.png.", antes=14, tam=9.5, cursiva=True, color=GRIS)

# ------------------------------------------------------- numero de pagina --
# El campo va en la primera seccion; las demas heredan el mismo pie.
par_pie = doc.sections[0].footer.paragraphs[0]
par_pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = par_pie.add_run()
run.font.size = Pt(9)
run.font.color.rgb = GRIS
ini = OxmlElement("w:fldChar"); ini.set(qn("w:fldCharType"), "begin")
instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
instr.text = " PAGE "
fin = OxmlElement("w:fldChar"); fin.set(qn("w:fldCharType"), "end")
for nodo in (ini, instr, fin):
    run._r.append(nodo)

for s in doc.sections[1:]:
    s.footer.is_linked_to_previous = True

os.makedirs("entregables", exist_ok=True)
ruta = os.path.join("entregables", "Proyecto_Integrador_TicketFlow.docx")
doc.save(ruta)
print("Listo:", ruta, "%.1f KB" % (os.path.getsize(ruta) / 1024))
