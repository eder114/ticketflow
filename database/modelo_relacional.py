# -*- coding: utf-8 -*-
"""
Genera el documento Word con el paso del Modelo Entidad-Relación de
TicketFlow al Modelo Relacional, siguiendo los pasos de la asignatura
(Tema 3 · Modelo Relacional y el documento complementario de pasos de
transformación).

Salida: database/Modelo_Relacional_TicketFlow.docx
"""

import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

AQUI = os.path.dirname(os.path.abspath(__file__))
SALIDA = os.path.join(AQUI, "Modelo_Relacional_TicketFlow.docx")

AZUL = RGBColor(0x49, 0x3E, 0xE5)
TINTA = RGBColor(0x19, 0x1C, 0x1E)
GRIS = RGBColor(0x5A, 0x58, 0x68)
ROJO = RGBColor(0xB4, 0x23, 0x2A)

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"; st.font.size = Pt(11); st.font.color.rgb = TINTA
st._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.2
for nivel, tam in ((1, 16), (2, 13)):
    h = doc.styles["Heading %d" % nivel]
    h.font.name = "Cambria"; h.font.size = Pt(tam); h.font.bold = True
    h.font.color.rgb = TINTA if nivel == 1 else AZUL
    h.paragraph_format.space_before = Pt(15 if nivel == 1 else 10)
    h.paragraph_format.space_after = Pt(6); h.paragraph_format.keep_with_next = True
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(2.2)
    s.left_margin = s.right_margin = Cm(2.3)
ANCHO = 16.4


def p(texto="", tam=None, negrita=False, cursiva=False, color=None, centro=False, antes=None, despues=None):
    par = doc.add_paragraph()
    if texto:
        r = par.add_run(texto); r.bold = negrita; r.italic = cursiva
        if tam: r.font.size = Pt(tam)
        if color: r.font.color.rgb = color
    if centro: par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if antes is not None: par.paragraph_format.space_before = Pt(antes)
    if despues is not None: par.paragraph_format.space_after = Pt(despues)
    return par


def vineta(texto, negrita_inicial=""):
    par = doc.add_paragraph(style="List Bullet")
    par.paragraph_format.space_after = Pt(3)
    if negrita_inicial:
        par.add_run(negrita_inicial).bold = True
    par.add_run(texto)
    return par


def sombra(celda, color):
    tc = celda._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), color); tc.append(sh)


def tabla(cab, filas, anchos, tam=9.5, centrar=()):
    t = doc.add_table(rows=1, cols=len(cab)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(cab):
        c = t.rows[0].cells[i]; c.text = ""
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(tam); r.font.color.rgb = GRIS
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER; sombra(c, "EDECFC")
    for fila in filas:
        cs = t.add_row().cells
        for i, v in enumerate(fila):
            cs[i].text = ""; par = cs[i].paragraphs[0]; par.paragraph_format.space_after = Pt(1)
            r = par.add_run(str(v)); r.font.size = Pt(tam)
            if i in centrar: par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for fila in t.rows:
        for i, a in enumerate(anchos):
            fila.cells[i].width = Cm(a)
    return t


def esquema(nombre, atributos, sangria=0.6):
    """Escribe  NOMBRE(pk, atributo, fk)  con la PK subrayada y la FK marcada.

    atributos: lista de (texto, tipo) con tipo "" | "pk" | "fk" | "pkfk"
    """
    par = doc.add_paragraph(); par.paragraph_format.left_indent = Cm(sangria)
    par.paragraph_format.space_after = Pt(2)
    r = par.add_run(nombre); r.bold = True; r.font.name = "Consolas"; r.font.size = Pt(10.5)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    r = par.add_run(" ("); r.font.name = "Consolas"; r.font.size = Pt(10.5)
    for i, (txt, tipo) in enumerate(atributos):
        if i:
            s2 = par.add_run(", "); s2.font.name = "Consolas"; s2.font.size = Pt(10.5)
        r = par.add_run(txt)
        r.font.name = "Consolas"; r.font.size = Pt(10.5)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        if tipo in ("pk", "pkfk"):
            r.underline = True; r.bold = True
        if tipo in ("fk", "pkfk"):
            r.font.color.rgb = AZUL
            s3 = par.add_run("*"); s3.font.name = "Consolas"; s3.font.size = Pt(10.5); s3.font.color.rgb = AZUL
    r = par.add_run(")"); r.font.name = "Consolas"; r.font.size = Pt(10.5)
    return par


def nota_fk(texto):
    par = doc.add_paragraph(); par.paragraph_format.left_indent = Cm(1.2)
    par.paragraph_format.space_after = Pt(8)
    r = par.add_run(texto); r.font.size = Pt(9); r.font.color.rgb = GRIS; r.italic = True
    return par


def sql(bloque):
    par = doc.add_paragraph(); par.paragraph_format.left_indent = Cm(0.4)
    par.paragraph_format.space_after = Pt(10)
    r = par.add_run(bloque); r.font.name = "Consolas"; r.font.size = Pt(9)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    return par


def salto():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# =========================================================== portada ======
for _ in range(4):
    doc.add_paragraph()
p("TICKETFLOW", tam=28, negrita=True, color=AZUL, centro=True, despues=2)
p("Del Modelo Entidad-Relación al Modelo Relacional", tam=16, negrita=True, centro=True, despues=4)
p("Transformación paso a paso del diagrama E-R del proyecto", tam=12, color=GRIS, centro=True, despues=28)
p("Bases de Datos y Programación en Ambiente Web I · Proyecto Integrador 2026-2",
  tam=10.5, color=GRIS, centro=True, despues=24)
for n in ("Eder Fabián Rodríguez Murillo", "Eduardo José Benítez Guevara",
          "Jorge Andrés Marín Díaz", "Samuel Uribe Naranjo"):
    p(n, tam=11.5, centro=True, despues=2)
p("Unidad Central del Valle del Cauca — UCEVA · Tuluá", tam=11, negrita=True, centro=True, antes=28, despues=2)
p("Septiembre de 2026", tam=11, color=GRIS, centro=True)
salto()

# =========================================================== 1. punto de partida
doc.add_heading("1. Punto de partida: el diagrama Entidad-Relación", level=1)
p("El siguiente diagrama, en notación de Chen, es el modelo conceptual del proyecto. "
  "De él salen las diez relaciones (tablas) del modelo relacional.")
img = os.path.join(AQUI, "DER_Chen_TicketFlow.png")
if os.path.exists(img):
    doc.add_picture(img, width=Cm(ANCHO))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p("Figura 1. Diagrama Entidad-Relación de TicketFlow (notación de Chen).",
      tam=9, cursiva=True, color=GRIS, centro=True, despues=10)

doc.add_heading("Notación que se usa en este documento", level=2)
vineta("cada relación se escribe como NOMBRE (atributo1, atributo2, …).", "Esquema: ")
vineta("va subrayada y en negrita.", "Llave primaria (PK): ")
vineta("va en azul y marcada con un asterisco (*). Debajo de cada relación se indica a qué tabla apunta.", "Llave foránea (FK): ")
vineta("una llave que además es foránea aparece subrayada y en azul.", "PK que también es FK: ")
salto()

# =========================================================== 2. pasos =====
doc.add_heading("2. Transformación paso a paso", level=1)

# --- Paso 1
doc.add_heading("Paso 1 · Entidades fuertes", level=2)
p("Regla: por cada entidad fuerte se crea una relación con una columna por cada atributo "
  "simple, y el atributo identificador pasa a ser la llave primaria.")
p("En nuestro diagrama son fuertes PAIS, DEPARTAMENTO, CIUDAD, PERSONA, EVENTO y RESERVA. "
  "El atributo compuesto nombre_completo se descompone en sus partes simples (nombres y "
  "apellidos), como indica el paso 2 del documento de transformación.")
esquema("PAIS", [("id_pais", "pk"), ("nombre", "")])
esquema("DEPARTAMENTO", [("id_departamento", "pk"), ("nombre", "")])
esquema("CIUDAD", [("id_ciudad", "pk"), ("nombre", "")])
esquema("PERSONA", [("identificacion", "pk"), ("nombres", ""), ("apellidos", ""), ("correo", ""), ("direccion", "")])
esquema("EVENTO", [("codigo_evento", "pk"), ("nombre", ""), ("descripcion", ""), ("teatro", ""),
                   ("fecha_hora_inicio", ""), ("fecha_hora_fin", ""), ("capacidad_total", ""),
                   ("precio_base", ""), ("observaciones", ""), ("estado", "")])
esquema("RESERVA", [("id_reserva", "pk"), ("fecha_hora", ""), ("numero_entradas", ""),
                    ("observaciones", ""), ("estado", "")])
p("Los atributos derivados cupos_disponibles y valor_total no se crean como columnas: se "
  "calculan cuando se consultan, para que no queden dos datos que puedan contradecirse.",
  tam=10, cursiva=True, color=GRIS)

# --- Paso 2
doc.add_heading("Paso 2 · Entidad débil", level=2)
p("Regla: la entidad débil se convierte en una relación que incluye como llave foránea la "
  "llave primaria de su entidad fuerte. La llave primaria se forma con esa llave foránea "
  "más el identificador parcial.")
p("TELEFONO es débil y depende de PERSONA: un número por sí solo no identifica nada, pero "
  "junto con la identificación de su dueño sí.")
esquema("TELEFONO", [("identificacion", "pkfk"), ("numero", "pk")])
nota_fk("FK: identificacion → PERSONA(identificacion). PK compuesta: (identificacion, numero).")

# --- Paso 3
doc.add_heading("Paso 3 · Relaciones 1:N", level=2)
p("Regla: la llave primaria del lado 1 pasa como llave foránea a la relación del lado N. "
  "La relación desaparece como tabla, porque no tiene atributos propios.")
tabla(["Relación del diagrama", "Lado 1", "Lado N", "Llave foránea que se agrega"],
      [["tiene", "PAIS", "DEPARTAMENTO", "DEPARTAMENTO.id_pais"],
       ["tiene", "DEPARTAMENTO", "CIUDAD", "CIUDAD.id_departamento"],
       ["reside en", "CIUDAD", "PERSONA", "PERSONA.id_ciudad"],
       ["se realiza en", "CIUDAD", "EVENTO", "EVENTO.id_ciudad"],
       ["registra", "AGENTE", "EVENTO", "EVENTO.identificacion_agente"],
       ["realiza", "CLIENTE", "RESERVA", "RESERVA.identificacion_cliente"],
       ["genera", "EVENTO", "RESERVA", "RESERVA.codigo_evento"]],
      [4.2, 3.4, 3.4, 5.4])
p("")
esquema("DEPARTAMENTO", [("id_departamento", "pk"), ("nombre", ""), ("id_pais", "fk")])
nota_fk("FK: id_pais → PAIS(id_pais)")
esquema("CIUDAD", [("id_ciudad", "pk"), ("nombre", ""), ("id_departamento", "fk")])
nota_fk("FK: id_departamento → DEPARTAMENTO(id_departamento)")
esquema("PERSONA", [("identificacion", "pk"), ("nombres", ""), ("apellidos", ""), ("correo", ""),
                    ("direccion", ""), ("id_ciudad", "fk")])
nota_fk("FK: id_ciudad → CIUDAD(id_ciudad)")

# --- Paso 4
doc.add_heading("Paso 4 · Relaciones 1:1 y jerarquía de subtipos", level=2)
p("CLIENTE, AGENTE y ADMINISTRADOR son subtipos de PERSONA, unidos por relaciones 1:1. "
  "El documento de transformación plantea tres caminos para una jerarquía: reunir todo en "
  "una sola tabla, crear una tabla por subtipo eliminando el supertipo, o crear una tabla "
  "para el supertipo y una por cada subtipo.")
p("Elegimos el tercero, que es el más usado y el que pide el enunciado del proyecto, porque "
  "este exige registrar, modificar, eliminar y consultar por separado personas, clientes, "
  "agentes y administradores. Cada subtipo recibe como llave primaria la llave del supertipo, "
  "que al mismo tiempo es llave foránea.")
esquema("CLIENTE", [("identificacion", "pkfk"), ("puntos", ""), ("ve_publicidad", "")])
nota_fk("FK: identificacion → PERSONA(identificacion)")
esquema("AGENTE", [("identificacion", "pkfk"), ("comision", ""), ("experiencia", "")])
nota_fk("FK: identificacion → PERSONA(identificacion)")
esquema("ADMINISTRADOR", [("identificacion", "pkfk"), ("salario", ""), ("horario", "")])
nota_fk("FK: identificacion → PERSONA(identificacion)")

# --- Paso 5
doc.add_heading("Paso 5 · Relaciones N:M", level=2)
p("Regla: una relación N:M se convierte en una tabla nueva, cuya llave primaria es la unión "
  "de las llaves de las entidades que asocia.")
p("En el diagrama no queda ninguna N:M sin resolver. Entre CLIENTE y EVENTO existiría una "
  "(un cliente asiste a muchos eventos y un evento recibe muchos clientes), pero ya la "
  "resolvimos en el modelo conceptual: RESERVA es una entidad propia, con fecha, número de "
  "entradas, valor y estado, y por eso se relaciona 1:N con cada una de las dos.")
p("Si RESERVA no guardara datos propios, la transformación habría sido:", tam=10, cursiva=True, color=GRIS)
esquema("RESERVA_N_M", [("identificacion_cliente", "pkfk"), ("codigo_evento", "pkfk")])
nota_fk("Esta es la forma que NO usamos: perdería la fecha, el valor y el estado de la reserva.")

# --- Paso 6
doc.add_heading("Paso 6 · Atributos multivaluados", level=2)
p("Regla: un atributo multivaluado se convierte en una tabla aparte, con la llave foránea de "
  "su entidad y una columna para el atributo.")
p("El teléfono es el único caso: el enunciado dice que una persona puede tener varios. Lo "
  "modelamos como entidad débil en el paso 2, y el resultado es exactamente el mismo que si "
  "lo hubiéramos tratado como atributo multivaluado: la tabla TELEFONO con llave primaria "
  "(identificacion, numero).")

# --- Paso 7
doc.add_heading("Paso 7 · Relaciones n-arias", level=2)
p("Regla: una relación en la que participan más de dos entidades se convierte en una tabla "
  "con las llaves foráneas de todas ellas.")
p("No aplica: las once relaciones del diagrama son binarias. La que podría parecer ternaria "
  "—cliente, evento y reserva— no lo es, porque la reserva es una entidad y no una relación.")
salto()

# =========================================================== 3. esquema final
doc.add_heading("3. Esquema relacional resultante", level=1)
p("Diez relaciones. La llave primaria va subrayada y las llaves foráneas en azul con asterisco.")

esquema("PAIS", [("id_pais", "pk"), ("nombre", "")])
esquema("DEPARTAMENTO", [("id_departamento", "pk"), ("nombre", ""), ("id_pais", "fk")])
esquema("CIUDAD", [("id_ciudad", "pk"), ("nombre", ""), ("id_departamento", "fk")])
esquema("PERSONA", [("identificacion", "pk"), ("nombres", ""), ("apellidos", ""), ("correo", ""),
                    ("direccion", ""), ("id_ciudad", "fk")])
esquema("TELEFONO", [("identificacion", "pkfk"), ("numero", "pk")])
esquema("CLIENTE", [("identificacion", "pkfk"), ("puntos", ""), ("ve_publicidad", "")])
esquema("AGENTE", [("identificacion", "pkfk"), ("comision", ""), ("experiencia", "")])
esquema("ADMINISTRADOR", [("identificacion", "pkfk"), ("salario", ""), ("horario", "")])
esquema("EVENTO", [("codigo_evento", "pk"), ("nombre", ""), ("descripcion", ""), ("teatro", ""),
                   ("fecha_hora_inicio", ""), ("fecha_hora_fin", ""), ("capacidad_total", ""),
                   ("precio_base", ""), ("observaciones", ""), ("estado", ""),
                   ("id_ciudad", "fk"), ("identificacion_agente", "fk")])
esquema("RESERVA", [("id_reserva", "pk"), ("fecha_hora", ""), ("numero_entradas", ""),
                    ("observaciones", ""), ("estado", ""), ("identificacion_cliente", "fk"),
                    ("codigo_evento", "fk")])

doc.add_heading("Llaves foráneas", level=2)
tabla(["Tabla", "Llave foránea", "Apunta a", "Qué representa"],
      [["DEPARTAMENTO", "id_pais", "PAIS(id_pais)", "El país del departamento"],
       ["CIUDAD", "id_departamento", "DEPARTAMENTO(id_departamento)", "El departamento de la ciudad"],
       ["PERSONA", "id_ciudad", "CIUDAD(id_ciudad)", "La ciudad donde reside"],
       ["TELEFONO", "identificacion", "PERSONA(identificacion)", "El dueño del teléfono"],
       ["CLIENTE", "identificacion", "PERSONA(identificacion)", "La persona que es cliente"],
       ["AGENTE", "identificacion", "PERSONA(identificacion)", "La persona que es agente"],
       ["ADMINISTRADOR", "identificacion", "PERSONA(identificacion)", "La persona que es administrador"],
       ["EVENTO", "id_ciudad", "CIUDAD(id_ciudad)", "La ciudad donde se realiza"],
       ["EVENTO", "identificacion_agente", "AGENTE(identificacion)", "El agente que lo registró"],
       ["RESERVA", "identificacion_cliente", "CLIENTE(identificacion)", "El cliente que reserva"],
       ["RESERVA", "codigo_evento", "EVENTO(codigo_evento)", "El evento reservado"]],
      [3.4, 4.0, 5.0, 4.0])

doc.add_heading("Restricciones de integridad", level=2)
vineta("ninguna llave primaria admite nulos, y no hay dos filas con la misma llave.", "Integridad de entidad: ")
vineta("toda llave foránea debe existir en la tabla a la que apunta. Un teléfono sin persona o una reserva sin evento no se pueden guardar.", "Integridad referencial: ")
vineta("estado de EVENTO solo acepta Programado, En Boletería, En Vivo, Finalizado o Cancelado; estado de RESERVA solo Reservada, Confirmada o Cancelada.", "Dominio: ")
vineta("el correo de PERSONA no se puede repetir.", "Unicidad: ")
vineta("capacidad_total y numero_entradas deben ser mayores que cero, precio_base no puede ser negativo y la fecha de fin debe ser posterior a la de inicio.", "Valores válidos: ")
vineta("las entradas reservadas y confirmadas de un evento no pueden superar su capacidad_total. Esta se revisa en el momento de confirmar la reserva.", "Regla del negocio: ")
salto()

# =========================================================== 4. SQL =======
doc.add_heading("4. Anexo: el esquema en PostgreSQL", level=1)
p("Traducción directa del esquema anterior al lenguaje de definición de datos, que es lo que "
  "se implementará en el Sprint 2.")
sql("""CREATE TABLE pais (
    id_pais        SERIAL PRIMARY KEY,
    nombre         VARCHAR(60) NOT NULL
);

CREATE TABLE departamento (
    id_departamento SERIAL PRIMARY KEY,
    nombre          VARCHAR(60) NOT NULL,
    id_pais         INTEGER NOT NULL REFERENCES pais(id_pais)
);

CREATE TABLE ciudad (
    id_ciudad       SERIAL PRIMARY KEY,
    nombre          VARCHAR(60) NOT NULL,
    id_departamento INTEGER NOT NULL REFERENCES departamento(id_departamento)
);

CREATE TABLE persona (
    identificacion  VARCHAR(20) PRIMARY KEY,
    nombres         VARCHAR(60) NOT NULL,
    apellidos       VARCHAR(60) NOT NULL,
    correo          VARCHAR(120) NOT NULL UNIQUE,
    direccion       VARCHAR(120),
    id_ciudad       INTEGER NOT NULL REFERENCES ciudad(id_ciudad)
);

CREATE TABLE telefono (
    identificacion  VARCHAR(20) REFERENCES persona(identificacion),
    numero          VARCHAR(20),
    PRIMARY KEY (identificacion, numero)
);

CREATE TABLE cliente (
    identificacion  VARCHAR(20) PRIMARY KEY REFERENCES persona(identificacion),
    puntos          INTEGER NOT NULL DEFAULT 0,
    ve_publicidad   BOOLEAN NOT NULL DEFAULT TRUE
);

CREATE TABLE agente (
    identificacion  VARCHAR(20) PRIMARY KEY REFERENCES persona(identificacion),
    comision        NUMERIC(5,2) NOT NULL DEFAULT 0,
    experiencia     INTEGER NOT NULL DEFAULT 0
);

CREATE TABLE administrador (
    identificacion  VARCHAR(20) PRIMARY KEY REFERENCES persona(identificacion),
    salario         NUMERIC(12,2) NOT NULL,
    horario         VARCHAR(60)
);

CREATE TABLE evento (
    codigo_evento     VARCHAR(20) PRIMARY KEY,
    nombre            VARCHAR(120) NOT NULL,
    descripcion       TEXT,
    teatro            VARCHAR(120),
    fecha_hora_inicio TIMESTAMP NOT NULL,
    fecha_hora_fin    TIMESTAMP,
    capacidad_total   INTEGER NOT NULL CHECK (capacidad_total > 0),
    precio_base       NUMERIC(12,2) NOT NULL CHECK (precio_base >= 0),
    observaciones     TEXT,
    estado            VARCHAR(20) NOT NULL
                      CHECK (estado IN ('Programado','En Boleteria','En Vivo',
                                        'Finalizado','Cancelado')),
    id_ciudad         INTEGER NOT NULL REFERENCES ciudad(id_ciudad),
    identificacion_agente VARCHAR(20) NOT NULL REFERENCES agente(identificacion),
    CHECK (fecha_hora_fin IS NULL OR fecha_hora_fin > fecha_hora_inicio)
);

CREATE TABLE reserva (
    id_reserva      SERIAL PRIMARY KEY,
    fecha_hora      TIMESTAMP NOT NULL DEFAULT NOW(),
    numero_entradas INTEGER NOT NULL CHECK (numero_entradas > 0),
    observaciones   TEXT,
    estado          VARCHAR(20) NOT NULL
                    CHECK (estado IN ('Reservada','Confirmada','Cancelada')),
    identificacion_cliente VARCHAR(20) NOT NULL REFERENCES cliente(identificacion),
    codigo_evento   VARCHAR(20) NOT NULL REFERENCES evento(codigo_evento)
);""")

doc.add_heading("Los atributos derivados", level=2)
p("No son columnas: se calculan con una consulta, para que nunca queden desactualizados.")
sql("""-- Cupos disponibles de un evento
SELECT e.capacidad_total - COALESCE(SUM(r.numero_entradas), 0) AS cupos_disponibles
FROM   evento e
LEFT   JOIN reserva r ON r.codigo_evento = e.codigo_evento
       AND r.estado IN ('Reservada','Confirmada')
WHERE  e.codigo_evento = 'EV-001'
GROUP  BY e.capacidad_total;

-- Valor total de una reserva
SELECT r.numero_entradas * e.precio_base AS valor_total
FROM   reserva r
JOIN   evento e ON e.codigo_evento = r.codigo_evento
WHERE  r.id_reserva = 1;""")

# ------------------------------------------------------- número de página --
par_pie = doc.sections[0].footer.paragraphs[0]
par_pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = par_pie.add_run(); run.font.size = Pt(9); run.font.color.rgb = GRIS
ini = OxmlElement("w:fldChar"); ini.set(qn("w:fldCharType"), "begin")
instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
fin = OxmlElement("w:fldChar"); fin.set(qn("w:fldCharType"), "end")
for nodo in (ini, instr, fin):
    run._r.append(nodo)

try:
    doc.save(SALIDA)
    print("Listo:", SALIDA, "%.1f KB" % (os.path.getsize(SALIDA) / 1024))
except PermissionError:
    print("OJO: el Word está abierto. Ciérrelo y vuelva a correr el script.")
